from runtime_paths import APP_DIR
import json,re,sqlite3,uuid,threading,zipfile
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from contextlib import contextmanager
from concurrent.futures import ThreadPoolExecutor,as_completed
from flask import Flask,request,jsonify,render_template,send_file,abort
from openpyxl import load_workbook
from cities import load_cities,save_cities,REGIONS
from recommendations import fetch_weather,recommend,CATEGORIES
from scraper import check_city
from exports import export_report
from analysis_context import capture_context,resolve_context
from gemini_recommendations import enrich
from recommendation_display import display_fields
ROOT=APP_DIR;DB=ROOT/'data'/'runs.sqlite3';DB.parent.mkdir(exist_ok=True)
app=Flask(__name__);app.config['MAX_CONTENT_LENGTH']=10*1024*1024
RUNS={};LOCK=threading.Lock()
@contextmanager
def connection():
 db=sqlite3.connect(DB,timeout=15);db.row_factory=sqlite3.Row
 db.execute('CREATE TABLE IF NOT EXISTS runs (id TEXT PRIMARY KEY,status TEXT,created TEXT,config TEXT,error TEXT)')
 db.execute('CREATE TABLE IF NOT EXISTS results (run_id TEXT,asin TEXT,pincode TEXT,result TEXT,advice TEXT,PRIMARY KEY(run_id,asin,pincode))')
 try:
  with db: yield db
 finally: db.close()

@app.before_request
def same_origin():
 if request.method not in ('GET','HEAD','OPTIONS'):
  origin=request.headers.get('Origin')
  if origin and origin.rstrip('/')!=request.host_url.rstrip('/'): abort(403)
def parse_products(text):
 values=[];seen=set()
 for line in str(text).splitlines():
  for value in re.split(r'[,;\s]+',line.strip()):
   if not value: continue
   value=value.upper()
   if not re.fullmatch(r'[A-Z0-9]{10}',value): raise ValueError('Every ASIN must have 10 letters/digits. Invalid: '+value[:30])
   if value not in seen: seen.add(value);values.append(dict(asin=value,brand=''))
 return values
@app.get('/')
def home(): return render_template('index.html')
@app.get('/api/cities')
def cities(): return jsonify(cities=load_cities(),regions=list(REGIONS))
@app.put('/api/cities')
def edit_cities():
 try: save_cities(request.get_json()['cities']);return jsonify(success=True)
 except (ValueError,TypeError,KeyError) as exc: return jsonify(message=str(exc)),400
@app.post('/api/import')
def import_asins():
 file=request.files.get('file')
 if not file or not file.filename.lower().endswith('.xlsx'): return jsonify(message='Choose an .xlsx workbook.'),400
 try:
  with zipfile.ZipFile(file.stream) as archive:
   if sum(entry.file_size for entry in archive.infolist())>100*1024*1024: raise ValueError('Workbook expands beyond 100 MB.')
  file.stream.seek(0);book=load_workbook(file.stream,read_only=True,data_only=True);products={}
  for sheet in book:
   iterator=sheet.iter_rows(values_only=True);header=next(iterator,());names=[str(value or '').strip().lower() for value in header]
   if 'asin' not in names: continue
   index=names.index('asin');brand=names.index('brand') if 'brand' in names else None
   for row in iterator:
    value=str(row[index] or '').strip().upper() if index<len(row) else ''
    if re.fullmatch(r'[A-Z0-9]{10}',value): products.setdefault(value,dict(asin=value,brand=str(row[brand] or '') if brand is not None and brand<len(row) else ''))
    if len(products)>1000: raise ValueError('Import up to 1,000 distinct ASINs.')
  book.close()
  if not products: raise ValueError('No ASIN column or valid ASINs found.')
  return jsonify(products=list(products.values()))
 except Exception as exc: return jsonify(message='Could not import workbook: '+str(exc)[:180]),400
@app.post('/api/runs')
def start_run():
 try:
  body=request.get_json() or {};products=parse_products(body.get('asins',''));ids=body.get('cities',[]);all_cities=load_cities();selected=[city for city in all_cities if city['id'] in ids]
  origin='Selected destination cities';category='auto'
  if not 1<=len(products)<=1000 or not selected or len(products)*len(selected)>10000: raise ValueError('Choose ASINs and cities, up to 10,000 checks per run.')
  if category not in CATEGORIES: raise ValueError('Choose a valid product category.')
  if body.get('show_browser') not in (None,True,False): raise ValueError('Choose a valid browser option.')
  brands={p.get('asin'):str(p.get('brand',''))[:100] for p in body.get('imported',[]) if isinstance(p,dict)}
  for product in products: product['brand']=brands.get(product['asin'],'')
  config=dict(analysis_context=capture_context(body,all_cities),products=products,cities=selected,origin=origin,location_mode='selected_destinations',category=category,show_browser=body.get('show_browser',True));run_id=uuid.uuid4().hex;created=datetime.now(ZoneInfo('Asia/Kolkata')).isoformat()
  with LOCK:
   if RUNS: return jsonify(message='A check is already running. Cancel it or wait for it to finish.'),409
   event=threading.Event();RUNS[run_id]=event
  with connection() as db: db.execute('INSERT INTO runs VALUES (?,?,?,?,?)',(run_id,'running',created,json.dumps(config),''))
  threading.Thread(target=run_job,args=(run_id,config,event),daemon=True).start();return jsonify(id=run_id),202
 except (ValueError,TypeError) as exc: return jsonify(message=str(exc)),400
def run_job(run_id,config,cancel):
 errors=[];blocked=[False]
 config['analysis_context']=resolve_context(config.get('analysis_context',capture_context({},[])))
 with connection() as db: db.execute('UPDATE runs SET config=? WHERE id=?',(json.dumps(config),run_id))
 def worker(city):
  if cancel.is_set(): return
  weather=fetch_weather(city)
  def save(result):
   advice=recommend(result,city,weather,city['name'] if config.get('location_mode')=='selected_destinations' else config['origin'],config['category'],config.get('analysis_context'))
   with connection() as db: db.execute('INSERT OR REPLACE INTO results VALUES (?,?,?,?,?)',(run_id,result['asin'],result['pincode'],json.dumps(result),json.dumps(advice)))
   if result.get('status')=='Verification needed': blocked[0]=True;cancel.set()
  check_city(config['products'],city,cancel,save,config['show_browser'])
 try:
  with ThreadPoolExecutor(max_workers=2) as pool:
   pending=[pool.submit(worker,city) for city in config['cities']]
   for task in as_completed(pending):
    try: task.result()
    except Exception as exc: errors.append(str(exc)[:200])
  if not cancel.is_set():
   _,_,ai_rows=read_run(run_id)
   config['ai_status']=enrich(ai_rows,cancel)
   with connection() as db:
    for advice in ai_rows:
     db.execute('UPDATE results SET advice=? WHERE run_id=? AND asin=? AND pincode=?',(json.dumps(advice),run_id,advice['asin'],advice['pincode']))
    db.execute('UPDATE runs SET config=? WHERE id=?',(json.dumps(config),run_id))
  status='verification_needed' if blocked[0] else 'cancelled' if cancel.is_set() else 'failed' if errors else 'completed'
  with connection() as db:
   checked=db.execute('SELECT result FROM results WHERE run_id=?',(run_id,)).fetchall()
   if status=='completed' and any(json.loads(row['result']).get('status') in ('Error','Verification needed') for row in checked): status='completed_with_errors'
   db.execute('UPDATE runs SET status=?,error=? WHERE id=?',(status,'; '.join(errors),run_id))
 finally:
  with LOCK: RUNS.pop(run_id,None)
def read_run(run_id):
 with connection() as db:
  row=db.execute('SELECT * FROM runs WHERE id=?',(run_id,)).fetchone()
  if not row: abort(404)
  rows=db.execute('SELECT result,advice FROM results WHERE run_id=? ORDER BY rowid',(run_id,)).fetchall()
 job=dict(row);job['config']=json.loads(job['config']);results=[json.loads(row['result']) for row in rows];stored_advice=[json.loads(row['advice']) for row in rows]
 metadata={}
 for result in results:
  if result.get('title'): metadata[result['asin']]=dict(title=result['title'],brand=result.get('brand'))
 for result in results:
  for field,value in metadata.get(result['asin'],{}).items():
   if not result.get(field) and value: result[field]=value
 advice=[]
 for result,entry in zip(results,stored_advice):
  weather=dict(status='available' if entry.get('temperature') is not None and entry.get('humidity') is not None else 'unavailable',temperature=entry.get('temperature'),humidity=entry.get('humidity'),precipitation=entry.get('precipitation'),time=entry.get('weather_time'),source=entry.get('weather_source','Unavailable'))
  context=job['config'].get('analysis_context') or dict(analysis_date=entry.get('analysis_date',job['created'][:10]),month=entry.get('month'),source_location='Unavailable — not captured for this run',source_weather={})
  current=recommend(result,next(city for city in job['config']['cities'] if city['pincode']==entry['pincode']),weather,entry.get('origin',''),'auto',context)
  current.update(purchase_badge=result.get('purchase_badge'),purchase_count_floor=result.get('purchase_count_floor'))
  if entry.get('ai_engine','').startswith('Gemini '):
   for field in ('weather_priority','sell_advice','seasonal_reason','ai_engine','basis','seasonal_priority','target_month','ai_seasonal_reason','ai_general_reason','seasonal_suitability_score'): current[field]=entry.get(field,current.get(field))
  current.setdefault('ai_engine','Local rules')
  current.update(display_fields(current))
  advice.append(current)
 return job,results,advice

@app.get('/api/runs')
def history():
 with connection() as db: rows=db.execute('SELECT id,status,created FROM runs ORDER BY created DESC LIMIT 20').fetchall()
 return jsonify(runs=[dict(row) for row in rows])
@app.get('/api/runs/<run_id>')
def run_status(run_id):
 job,results,advice=read_run(run_id);job.update(results=results,recommendations=advice,completed=len(results),total=len(job['config']['products'])*len(job['config']['cities']));return jsonify(job)
@app.post('/api/runs/<run_id>/cancel')
def cancel_run(run_id):
 with LOCK:
  event=RUNS.get(run_id)
  if event: event.set()
 return jsonify(success=True)
@app.get('/api/runs/<run_id>/export')
def export(run_id):
 job,results,advice=read_run(run_id)
 if not results: return jsonify(message='No checked results to export yet.'),400
 return send_file(export_report(job,results,advice),as_attachment=True,download_name='Amazon_Delivery_'+job['created'][:10]+'_'+run_id[:6]+'.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
@app.errorhandler(413)
def large_file(error): return jsonify(message='Upload a workbook smaller than 10 MB.'),413
if __name__=='__main__':
 with connection() as db: db.execute("UPDATE runs SET status='interrupted',error='Server restarted. Saved partial results can be exported; start a new check for remaining data.' WHERE status='running'")
 import argparse
 parser=argparse.ArgumentParser();parser.add_argument('--host',default='127.0.0.1');parser.add_argument('--port',type=int,default=5055);args=parser.parse_args();app.run(host=args.host,port=args.port,threaded=True,debug=False)
