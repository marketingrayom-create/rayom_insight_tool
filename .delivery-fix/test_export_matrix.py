import unittest
from openpyxl import load_workbook
from exports import export_report

class MatrixTests(unittest.TestCase):
 def test_city_columns_keep_one_row_per_asin(self):
  cities=[dict(name='Delhi',pincode='110001'),dict(name='Jaipur',pincode='302001')]
  products=[dict(asin='B0H5MMJB4H',brand='OPPO'),dict(asin='B0C57Z55Z3',brand='Energizer')]
  rows=[dict(asin='B0H5MMJB4H',pincode='110001',delivery_days=5,delivery_date='2026-10-18',status='Available',price=2799),dict(asin='B0H5MMJB4H',pincode='302001',delivery_days=3,delivery_date='2026-10-16',status='Available',price=2700)]
  book=load_workbook(export_report({'config':dict(cities=cities,products=products)},rows,[]))
  sheet=book['City Comparison']
  self.assertEqual(sheet.max_row,3)
  self.assertEqual(sheet.max_column,10)
  self.assertEqual(sheet['C1'].value,'Days (Delhi)')
  self.assertEqual(sheet['G1'].value,'Days (Jaipur)')
  self.assertEqual(sheet['F2'].value,2799)
  self.assertEqual(sheet['J2'].value,2700)
  self.assertEqual(sheet['D2'].number_format,'dd-mm-yyyy')
  self.assertEqual(sheet['E3'].value,'Not checked')
  self.assertIsNone(sheet['F3'].value)
  self.assertIn('110001',sheet['E2'].comment.text)
  self.assertEqual(sheet.freeze_panes,'C2')

if __name__=='__main__': unittest.main()
