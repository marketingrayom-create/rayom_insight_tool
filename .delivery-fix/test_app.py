import unittest,tempfile,json,time,io
from pathlib import Path
from unittest.mock import patch
from datetime import date
from openpyxl import load_workbook
import app as module
from cities import defaults,save_cities
from scraper import parse_delivery_date
from recommendations import recommend,category_for
from exports import export_report
class Tests(unittest.TestCase):
 def setUp(self):
  self.directory=tempfile.TemporaryDirectory();self.addCleanup(self.directory.cleanup)
  self.db=patch.object(module,'DB',Path(self.directory.name)/'runs.sqlite3');self.db.start();self.addCleanup(self.db.stop)
  self.client=module.app.test_client()
 def test_city_groups_have_ten_cities_per_region_and_fixed_pincodes(self):
  values=self.client.get('/api/cities').json['cities']
  for region in module.REGIONS: 
   for tier in (1,2,3,4): self.assertEqual(sum(city['region']==region and city['tier']==tier for city in values),3 if tier==4 else 10)
  self.assertTrue(all(len(city['pincode'])==6 for city in values));self.assertEqual(len(set(city['id'] for city in values)),198)
 def test_asin_import_deduplicates_reference_excel(self):
  from openpyxl import Workbook
  wb=Workbook();wb.active.append(['Brand','ASIN']);wb.active.append(['Example','B0C57Z55Z3']);wb.active.append(['Example','B0C57Z55Z3']);stream=io.BytesIO();wb.save(stream);stream.seek(0)
  result=self.client.post('/api/import',data={'file':(stream,'test.xlsx')})
  self.assertEqual(result.status_code,200);self.assertEqual(len(result.json['products']),1)
 def test_invalid_requests_and_cross_origin_are_rejected(self):
  self.assertEqual(self.client.post('/api/runs',json={'asins':'bad'}).status_code,400)
  self.assertEqual(self.client.put('/api/cities',json={},headers={'Origin':'https://other.example'}).status_code,403)
 def test_delivery_parser_handles_today_year_boundary_and_missing_dates(self):
  self.assertEqual(parse_delivery_date('Tomorrow',date(2026,12,31)),('2027-01-01',1))
  self.assertEqual(parse_delivery_date('FREE delivery 2 January',date(2026,12,31)),('2027-01-02',2))
  self.assertEqual(parse_delivery_date('No promise',date(2026,10,12)),(None,None))
 def test_weather_rules_do_not_guess_unknown_category_or_aqi(self):
  city=defaults()[0];weather=dict(status='available',temperature=32,humidity=60,time='2026-10-12T12:00',source='test')
  row=dict(asin='B0C57Z55Z3',title='Ceiling fan')
  result=recommend(row,city,weather,'Jaipur');self.assertEqual(result['seasonal_rating'],'Favourable now')
  result=recommend(dict(row,title='Air purifier'),city,weather,'Jaipur');self.assertEqual(result['seasonal_rating'],'AQI data needed')
  result=recommend(row,city,dict(status='unavailable',source='test'),'Jaipur');self.assertEqual(result['seasonal_rating'],'Weather unavailable')
  self.assertEqual(category_for('Unknown'),'general')
 def test_export_has_recommendation_sheet_dates_and_source_context(self):
  city=defaults()[0];job=dict(id='test',status='completed',created='2026-10-12T12:00:00+05:30',config=dict(cities=[city],products=[dict(asin='B0C57Z55Z3')],origin='Jaipur'))
  row=dict(asin='B0C57Z55Z3',pincode=city['pincode'],city=city['name'],status='Available',title='=unsafe',delivery_date='2026-10-14',delivery_days=2,price=100,location_verified=True)
  advice=recommend(row,city,dict(status='unavailable',source='test'),'Jaipur','fan')
  book=load_workbook(export_report(job,[row],[advice]));self.assertEqual(book.sheetnames,['City Comparison','Recommendations']);self.assertIn('city-specific sales are unavailable',book['Recommendations']['C2'].value);self.assertTrue(book['City Comparison']['C2'].value.startswith("'="));self.assertEqual(book['City Comparison']['G2'].value.date(),date(2026,10,14))
 def test_job_reports_verification_failure_and_keeps_partial_export(self):
  city=defaults()[0]
  def check(products,destination,cancel,callback,show):
   callback(dict(asin=products[0]['asin'],city=destination['name'],pincode=destination['pincode'],status='Verification needed',location_verified=False))
  with patch.object(module,'check_city',side_effect=check),patch.object(module,'fetch_weather',return_value=dict(status='unavailable',source='test')):
   result=self.client.post('/api/runs',json=dict(asins='B0C57Z55Z3',cities=[city['id']],origin='Jaipur',category='fan'))
   self.assertEqual(result.status_code,202)
   for _ in range(100):
    data=self.client.get('/api/runs/'+result.json['id']).json
    if data['status']!='running': break
    time.sleep(.02)
   self.assertEqual(data['status'],'verification_needed');self.assertEqual(data['completed'],1)
   self.assertEqual(self.client.get('/api/runs/'+result.json['id']+'/export').status_code,200)
if __name__=='__main__': unittest.main()

class SelectionTests(unittest.TestCase):
 setUp=Tests.setUp
 def test_selected_cities_do_not_require_a_typed_location(self):
  city=defaults()[0]
  with patch.object(module,'run_job'):
   result=self.client.post('/api/runs',json=dict(asins='B0D8HXX1H5',cities=[city['id']]))
   self.assertEqual(result.status_code,202)
   job=self.client.get('/api/runs/'+result.json['id']).json
   self.assertEqual([value['id'] for value in job['config']['cities']],[city['id']])
   self.assertEqual(job['config']['location_mode'],'selected_destinations')
   with module.LOCK: module.RUNS.pop(result.json['id'],None)

class BrowserTests(unittest.TestCase):
 def test_redirect_recovery_uses_loaded_page_without_repeat_navigation(self):
  from scraper import navigate
  class Page:
   main_frame='main';url='https://www.amazon.in/dp/B0D8HXX1H5'
   calls=0
   def on(self,event,callback): self.callback=callback
   def remove_listener(self,event,callback): pass
   def wait_for_load_state(self,*args,**kwargs): pass
   def goto(self,url,**kwargs):
    from types import SimpleNamespace
    self.calls+=1
    self.callback(SimpleNamespace(url=url,status=200,frame='main',request=SimpleNamespace(is_navigation_request=lambda:True)))
    raise RuntimeError('net::ERR_ABORTED')
  page=Page();response=navigate(page,page.url);self.assertEqual(response.status,200);self.assertEqual(page.calls,1)
 def test_aborted_wrong_page_is_not_accepted(self):
  from scraper import navigate
  from unittest.mock import Mock
  page=Mock();page.url='https://www.amazon.in/errors/validateCaptcha';page.goto.side_effect=RuntimeError('net::ERR_ABORTED')
  with self.assertRaises(RuntimeError): navigate(page,'https://www.amazon.in/dp/B0D8HXX1H5')
  page.goto.assert_called_once()

class DuplicateElementTests(unittest.TestCase):
 def test_duplicate_product_titles_use_visible_text_without_strict_mode_error(self):
  from playwright.sync_api import sync_playwright
  from scraper import visible_text
  with sync_playwright() as engine:
   browser=engine.chromium.launch(headless=True,executable_path=engine.chromium.executable_path)
   try:
    page=browser.new_page();page.set_content('<span id="productTitle" style="display:none">Hidden mobile title</span><span id="productTitle">The Better Home Stainless Steel 15L</span><a id="bylineInfo" style="display:none">Hidden brand</a><a id="bylineInfo">Visit the The Better Home Store</a>')
    self.assertEqual(visible_text(page,'#productTitle'),'The Better Home Stainless Steel 15L')
    self.assertEqual(visible_text(page,'#bylineInfo'),'Visit the The Better Home Store')
    self.assertIsNone(visible_text(page,'#unavailable'))
   finally: browser.close()

class RecommendationPriorityTests(unittest.TestCase):
 def test_fan_priority_uses_weather_not_guessed_sales_or_month(self):
  city=defaults()[0];context=dict(source_location='Jaipur',month='October',analysis_date='2026-10-12',source_weather={'temperature':22})
  result=recommend({'asin':'B0D8HXX1H5','title':'Ceiling fan'},city,dict(status='available',temperature=31,humidity=50,source='test'),'Jaipur',analysis_context=context)
  self.assertEqual(result['general_rating'],'Sales unavailable');self.assertEqual(result['sales_priority'],'Not ranked');self.assertEqual(result['weather_priority'],'High');self.assertIn('warmer',result['source_comparison']);self.assertEqual(result['month'],'October')
  result=recommend({'asin':'B0D8HXX1H5','title':'Ceiling fan'},city,dict(status='available',temperature=21,humidity=50,source='test'),'Jaipur',analysis_context=context)
  self.assertEqual(result['weather_priority'],'Low')
 def test_location_context_accepts_city_id_only_never_coordinates(self):
  from analysis_context import capture_context
  city=defaults()[0]
  result=capture_context(dict(source_city_id=city['id'],entered_at='2026-10-12T04:00:00Z',geolocation={'latitude':0,'longitude':0}),defaults())
  self.assertEqual(result['source_location'],city['name']);self.assertNotIn('coordinates',result);self.assertEqual(result['month'],'October')

class ComparisonCoverageTests(unittest.TestCase):
 def test_every_selected_city_is_exported_as_a_visible_row(self):
  selected=defaults()[:2];job=dict(id='test',created='2026-10-12T12:00:00+05:30',status='completed',config=dict(products=[{'asin':'B0DGPVXVY5'}],cities=selected))
  rows=[dict(asin='B0DGPVXVY5',title='Cricket bat',pincode=selected[0]['pincode'],status='Available',delivery_date='2026-10-21',delivery_days=9)]
  sheet=load_workbook(export_report(job,rows,[]))['City Comparison']
  self.assertEqual(sheet.max_row,3);self.assertEqual(sheet['D2'].value,selected[0]['name']);self.assertEqual(sheet['D3'].value,selected[1]['name']);self.assertEqual(sheet['H3'].value,'Not checked');self.assertIsNone(sheet['G3'].value)
 def test_cricket_classification_does_not_assume_missing_rain_is_zero(self):
  self.assertEqual(category_for('Strauss Clash Cricket Bat'),'sports')
  result=recommend({'asin':'B0DGPVXVY5','title':'Cricket bat'},defaults()[0],dict(status='available',temperature=25,humidity=50,source='test'),'Jaipur')
  self.assertEqual(result['weather_priority'],'Not ranked');self.assertIn('unavailable',result['seasonal_reason'])
