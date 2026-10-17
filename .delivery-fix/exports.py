from io import BytesIO
from recommendation_display import display_fields
from datetime import date,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

def safe(value):return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value

def export_report(job,results,recommendations):
 wb=Workbook();comparison=wb.active;comparison.title='City Comparison';advice=wb.create_sheet('Recommendations')
 matrix=len(job['config']['cities'])>1
 headers=['Brand','ASIN'] if matrix else ['Brand','ASIN','Product Title','City','Pincode','Days','Delivery Date','Status','Price (INR)','Error / Notes']
 if matrix:
  for city in job['config']['cities']:
   headers.extend([label+' ('+city['name']+')' for label in ('Days','Delivery Date','Status','Price')])
 comparison.append(headers)
 lookup={(row['asin'],row['pincode']):row for row in results};metadata={}
 for row in results:
  if row.get('title'):metadata[row['asin']]=row
 for product in job['config']['products']:
  info=metadata.get(product['asin'],{})
  if matrix:
   values=[safe(info.get('brand') or product.get('brand')),product['asin']]
   for city in job['config']['cities']:
    row=lookup.get((product['asin'],city['pincode']),{})
    values.extend([row.get('delivery_days'),date.fromisoformat(row['delivery_date']) if row.get('delivery_date') else None,row.get('status','Not checked'),row.get('price')])
   comparison.append(values)
   comparison.cell(comparison.max_row,2).comment=Comment('Product: '+str(info.get('title') or '')+'\nSource: https://www.amazon.in/dp/'+product['asin'],'Delivery Atlas')
  for city in job['config']['cities']:
   row=lookup.get((product['asin'],city['pincode']),{})
   if not matrix: comparison.append([safe(info.get('brand') or product.get('brand')),product['asin'],safe(info.get('title')),city['name'],city['pincode'],row.get('delivery_days'),date.fromisoformat(row['delivery_date']) if row.get('delivery_date') else None,row.get('status','Not checked'),row.get('price'),safe(row.get('error'))])
   status_column=5+4*job['config']['cities'].index(city) if matrix else 8
   comparison.cell(comparison.max_row,status_column).comment=Comment('City: '+city['name']+'\nPincode: '+city['pincode']+'\nNotes: '+str(row.get('error') or '')+'\nPromise: '+str(row.get('delivery_text') or 'Unavailable')+'\nPincode verified: '+str(row.get('location_verified',False))+'\nChecked: '+str(row.get('checked_at','Not checked'))+'\nSource: '+row.get('source_url','https://www.amazon.in/dp/'+product['asin']),'Delivery Atlas')
 advice.append(['ASIN','Product Name','General Recommendation','Seasonal Recommendation'])
 for row in sorted(recommendations,key=lambda r:(r['asin'],r['city'])):
  text=display_fields(row)
  advice.append([row['asin'],safe(row.get('title')),safe(text['general_recommendation']),safe(text['seasonal_recommendation'])])
  advice.cell(advice.max_row,4).comment=Comment('Engine: '+str(row.get('ai_engine','Local rules'))+'\n'+str(row.get('basis','')),'Delivery Atlas')
 widths={comparison:([18,16]+[14,23,24,18]*len(job['config']['cities']) if matrix else [16,15,36,17,11,8,16,21,14,35]),advice:[15,40,60,85]}
 for sheet in wb:
  sheet.freeze_panes=('C2' if matrix else 'D2') if sheet==comparison else 'C2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.zoomScale=85;sheet.row_dimensions[1].height=35
  sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
  for index,width in enumerate(widths[sheet],1):sheet.column_dimensions[get_column_letter(index)].width=width
  for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='244B7E' if sheet==comparison and matrix else '173F46');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(vertical='center',wrap_text=True)
  for row in sheet.iter_rows(min_row=2):
   for cell in row:
    cell.alignment=Alignment(vertical='top',wrap_text=True)
    if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F6F5')
    if isinstance(cell.value,date):cell.number_format='dd-mmm-yyyy'
   lines=max(len(str(cell.value or ''))//max(10,widths[sheet][cell.column-1]-3)+1 for cell in row);sheet.row_dimensions[row[0].row].height=22 if sheet==comparison and matrix else max(55,min(220,lines*15))
  if sheet==comparison:
   for row in sheet.iter_rows(min_row=2):
    if matrix:
     for offset in range(2,len(row),4):
      days,delivery,status,price=row[offset:offset+4]
      if isinstance(days.value,(int,float)):
       days.fill=PatternFill('solid',fgColor='C6EFCE' if days.value<=2 else 'D9E2F3' if days.value<=5 else 'FFEB9C')
      price.number_format='#,##0.00'
      status.fill=PatternFill('solid',fgColor='C6EFCE' if status.value=='Available' else 'FFF2CC' if status.value=='Not checked' else 'FFC7CE');status.font=Font(bold=True)
      delivery.number_format='dd-mm-yyyy'
    else:
     row[8].number_format='#,##0.00'
     row[7].fill=PatternFill('solid',fgColor='C6EFCE' if row[7].value=='Available' else 'FFC7CE');row[7].font=Font(bold=True)
 stream=BytesIO();wb.save(stream);stream.seek(0);return stream
